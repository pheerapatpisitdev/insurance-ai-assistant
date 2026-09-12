#!/usr/bin/env python3
"""Extract the ไอเฮลท์ตี้ อัลตร้า benefit table into data/riders/ihealthy-ultra.json.

The rate tables come from extract_rates.py; this reads what the rates cannot say — what
each of the six plans actually pays, and the contract terms that go on the page.

Deterministic: running twice on the same workbook yields byte-identical JSON.
Never edit data/riders/ihealthy-ultra.json by hand; edit this script and re-run
`npm run extract ihealthy-ultra`.
"""
import json
import re
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
XLSX = ROOT / "ไฟล์คำนวน" / "Sales proposal_Life Protect Plus 9 19 99_2026821_update benefit graph.xlsx"
OUT = ROOT / "data" / "riders" / "ihealthy-ultra.json"
BENEFIT_SHEET = "ผลประโยชน์ ไอเฮลท์ตี้ อัลตร้า"
TERMS_SHEET = "เงื่อนไขที่สำคัญ ข้อยกเว้น iHU"

# Benefit columns: the first of each plan's three (ผลประโยชน์ / จำนวนสูงสุด / ผลประโยชน์สูงสุด).
ADULT_COLS = {"SMART": "O", "BRONZE": "R", "SILVER": "U", "GOLD": "X", "DIAMOND": "AA", "PLATINUM": "AD"}
CHILD_COLS = {"SMART": "I", "BRONZE": "L"}
# The third column of each triple holds the annual maximum, merged down the whole block.
MAX_COLS = {"SMART": "Q", "BRONZE": "T", "SILVER": "W", "GOLD": "Z", "DIAMOND": "AC", "PLATINUM": "AF"}
PLAN_NAMES = {
    "SMART": "สมาร์ท", "BRONZE": "บรอนซ์", "SILVER": "ซิลเวอร์",
    "GOLD": "โกลด์", "DIAMOND": "ไดมอนด์", "PLATINUM": "แพลทินั่ม",
}
# extract_rates.py imports this to key the IHU rate table, so the plan numbers the two JSON
# files join on have one definition, here, rather than a copy in each script.
PLAN_NO = {"SMART": 1, "BRONZE": 2, "SILVER": 3, "GOLD": 4, "DIAMOND": 5, "PLATINUM": 6}
# The benefit sheet's own I4 ties the customer's participation to the plan number it reads from
# กรอกข้อมูล!R33: 1-2 → 30,000, 3-4 → 50,000, 5-6 → 100,000, or a co-payment of 20% instead.
DEDUCTIBLE = {1: 30_000, 2: 30_000, 3: 50_000, 4: 50_000, 5: 100_000, 6: 100_000}
COPAY_PERCENT = 20

# --- benefit sheet rows -----------------------------------------------------------------
FIRST_ROW, LAST_ROW = 7, 53
# Rows 39..53 are the บันทึกสลักหลัง (endorsement) block — benefits that ride on top of the
# rider itself, which ซิลเวอร์ and up buy and สมาร์ท/บรอนซ์ do not.
ENDORSEMENT_FROM = 39
# Rows that describe the customer's own choice rather than the plan, and the table headers.
SKIP_ROWS = {33, 34, 35, 36, 37, 38}
ANNUAL_MAX_ROW = 8
# The paragraphs printed under the table, and the disclaimer that closes the sheet.
BENEFIT_NOTE_ROWS = {
    "participationNote": 58,
    "noClaimDiscount": 62,
    "outOfTerritory": 63,
    "sharedLimit": 64,
    "premiumChanges": 67,
}
DISCLAIMER_ROW = 68

# --- terms sheet rows -------------------------------------------------------------------
TERMS_NOTE_ROWS = {"renewalCopay": 6, "exclusions": 25}
# The 120-day list runs down two newspaper columns, four rows deep.
DISEASE_ROWS, DISEASE_COLS = range(12, 16), ("A", "B")
PRE_EXISTING_ROWS = range(17, 23)
# The prose each hardcoded figure below was read from, asserted against it.
WAITING_ROW, SPECIAL_WAITING_ROW = 8, 10

# These five figures are the contract's own numbers. The workbook states them inside
# paragraphs rather than in cells of their own, so they are typed in here for the page to use
# as numbers — and each one is asserted to still appear in the paragraph it came from, so a
# re-worded workbook cannot leave the figure and its prose contradicting each other.
RENEWAL_TO_AGE = 98
WAITING_DAYS = 30
SPECIAL_WAITING_DAYS = 120
NO_CLAIM_DISCOUNT_PERCENT = 10
OUT_OF_TERRITORY_DAYS = 90

# What a correct read of this workbook produces; a re-shipped file that shifts a row or
# renames a sheet trips these rather than shipping a quietly shorter table.
EXPECTED_PLANS = 6
EXPECTED_ROWS = 41
EXPECTED_HEADINGS = 5
EXPECTED_DISEASES = 8
FIRST_SECTION_TITLE = "หมวดที่ 1"


def reader(ws):
    """Return text(col, row) for one sheet, with that sheet's own merged cells resolved.

    A read that does not resolve merges returns "" for every cell a merge covers — row 44 and
    all six annual maximums here — so the map is built once, with the sheet it belongs to, and
    no caller can hand it the wrong one.
    """
    merged = {}
    for rng in ws.merged_cells.ranges:
        anchor = ws.cell(rng.min_row, rng.min_col).coordinate
        for row in ws[rng.coord]:
            for c in row:
                merged[c.coordinate] = anchor

    def text(col, row):
        v = ws[merged.get(f"{col}{row}", f"{col}{row}")].value
        if v is None:
            return ""
        # Excel hands a whole number back as a float: cell U43 is 6000.0, and the page prints
        # what it is given, so it becomes 6000 before it becomes a string.
        if isinstance(v, float) and v == int(v):
            v = int(v)
        return re.sub(r"\s+", " ", str(v)).strip()

    return text


def leading_number(s):
    m = re.match(r"([\d,]+)", s)
    return int(m.group(1).replace(",", "")) if m else None


def section_no(title):
    """'หมวดที่ 14 ...' → 14. Sub-headings ('หมวดย่อยที่ 2.1') and group headings have none."""
    m = re.match(r"หมวดที่\s*(\d+)", title)
    return int(m.group(1)) if m else None


def extract():
    wb = openpyxl.load_workbook(XLSX, data_only=True)
    benefit = reader(wb[BENEFIT_SHEET])
    terms = reader(wb[TERMS_SHEET])

    plans = []
    for code, col in MAX_COLS.items():
        annual_max = benefit(col, ANNUAL_MAX_ROW)
        plans.append({
            "code": code,
            "name": PLAN_NAMES[code],
            "planNo": PLAN_NO[code],
            "annualMax": leading_number(annual_max),
            "annualMaxNote": annual_max,
            "deductible": DEDUCTIBLE[PLAN_NO[code]],
        })
    assert len(plans) == EXPECTED_PLANS, f"expected {EXPECTED_PLANS} plans, got {len(plans)}"
    no_max = [p["code"] for p in plans if not p["annualMax"]]
    assert not no_max, f"{BENEFIT_SHEET} row {ANNUAL_MAX_ROW}: no annual maximum for {no_max}"

    rows = []
    for r in range(FIRST_ROW, LAST_ROW + 1):
        if r in SKIP_ROWS:
            continue
        title = benefit("A", r)
        if not title:
            continue
        adult = {code: benefit(col, r) for code, col in ADULT_COLS.items()}
        if not any(adult.values()):
            rows.append({"heading": title})
            continue
        child = {code: benefit(col, r) for code, col in CHILD_COLS.items()}
        entry = {
            "no": section_no(title),
            "title": title,
            "endorsement": r >= ENDORSEMENT_FROM,
            "adult": adult,
        }
        # Only สมาร์ท and บรอนซ์ sell to a child, and only where the wording differs is the
        # child column worth carrying: everything else is the adult row.
        if any(child[c] != adult[c] for c in CHILD_COLS):
            entry["child"] = child
        rows.append(entry)

    headings = [r for r in rows if "heading" in r]
    benefit_rows = [r for r in rows if "adult" in r]
    assert len(rows) == EXPECTED_ROWS, f"expected {EXPECTED_ROWS} rows, got {len(rows)}"
    assert len(headings) == EXPECTED_HEADINGS, f"expected {EXPECTED_HEADINGS} headings, got {len(headings)}"
    assert benefit_rows[0]["title"].startswith(FIRST_SECTION_TITLE), \
        f"the table no longer starts at {FIRST_SECTION_TITLE}: {benefit_rows[0]['title'][:40]!r}"

    diseases = []
    for r in DISEASE_ROWS:
        for c in DISEASE_COLS:
            v = terms(c, r).lstrip("- ").strip()
            if v:
                diseases.append(v)
    assert len(diseases) == EXPECTED_DISEASES, f"expected {EXPECTED_DISEASES} illnesses, got {len(diseases)}"

    waiting_prose = terms("A", WAITING_ROW)
    special_waiting_prose = terms("A", SPECIAL_WAITING_ROW)
    notes = {key: benefit("A", r) for key, r in BENEFIT_NOTE_ROWS.items()}
    out_terms = {
        "renewalToAge": RENEWAL_TO_AGE,
        "waitingDays": WAITING_DAYS,
        "specialWaitingDays": SPECIAL_WAITING_DAYS,
        "specialWaitingDiseases": diseases,
        "noClaimDiscountPercent": NO_CLAIM_DISCOUNT_PERCENT,
        "outOfTerritoryDays": OUT_OF_TERRITORY_DAYS,
        "renewalCopay": terms("A", TERMS_NOTE_ROWS["renewalCopay"]),
        "preExisting": " ".join(t for t in (terms("A", r) for r in PRE_EXISTING_ROWS) if t),
        "exclusions": terms("A", TERMS_NOTE_ROWS["exclusions"]),
        "premiumChanges": notes["premiumChanges"],
        "outOfTerritory": notes["outOfTerritory"],
        "noClaimDiscount": notes["noClaimDiscount"],
        "sharedLimit": notes["sharedLimit"],
        "participationNote": notes["participationNote"],
    }
    empty = sorted(k for k, v in out_terms.items() if not v)
    assert not empty, f"nothing read for terms {empty} — check the row numbers at the top of this script"
    # Each figure is looked for in the words that carry it, not on its own: "9" sits inside
    # "98 ปี", so a digit lost off a constant would still be found in the prose it contradicts.
    for phrase, prose_name in (
        (f"ถึงอายุ {RENEWAL_TO_AGE} ปี", "renewalCopay"),
        (f"ร้อยละ {NO_CLAIM_DISCOUNT_PERCENT}", "noClaimDiscount"),
        (f"{OUT_OF_TERRITORY_DAYS} วัน", "outOfTerritory"),
    ):
        assert phrase in out_terms[prose_name], f"{phrase!r} is no longer stated in {prose_name}"
    assert f"{WAITING_DAYS} วัน" in waiting_prose, \
        f"'{WAITING_DAYS} วัน' is no longer stated in {TERMS_SHEET}!A{WAITING_ROW}"
    assert f"{SPECIAL_WAITING_DAYS} วัน" in special_waiting_prose, \
        f"'{SPECIAL_WAITING_DAYS} วัน' is no longer stated in {TERMS_SHEET}!A{SPECIAL_WAITING_ROW}"

    disclaimer = benefit("A", DISCLAIMER_ROW)
    assert "เอกสารประกอบการเสนอขาย" in disclaimer, \
        f"{BENEFIT_SHEET}!A{DISCLAIMER_ROW} is no longer the disclaimer: {disclaimer[:40]!r}"

    data = {
        "name": "ไอเฮลท์ตี้ อัลตร้า",
        "code": "IHU",
        "source": XLSX.name,
        "plans": plans,
        "copayPercent": COPAY_PERCENT,
        "rows": rows,
        "terms": out_terms,
        "disclaimer": disclaimer,
    }
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(data, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"wrote {OUT} ({OUT.stat().st_size} bytes)")


if __name__ == "__main__":
    extract()
