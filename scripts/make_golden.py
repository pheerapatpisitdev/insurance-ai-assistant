#!/usr/bin/env python3
"""Generate tests/golden/plb.json by letting LibreOffice recalculate the company Excel file.

For each case: copy the workbook, write inputs into 'กรอกข้อมูล', batch-convert with
`soffice --headless` (forced recalculation via a private profile), read back outputs.
"""
import json
import random
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "ไฟล์คำนวน" / "Protection Life (PLB)_A2026-1_01042026.xlsx"
CASES = ROOT / "tests" / "golden" / "cases-plb.json"
OUT = ROOT / "tests" / "golden" / "plb.json"
SOFFICE = "/Applications/LibreOffice.app/Contents/MacOS/soffice"

SEX_TH = {"M": "ชาย", "F": "หญิง"}
MODE_TH = {"annual": "รายปี", "semi": "ราย 6 เดือน", "monthly": "รายเดือน"}
VARIANTS = ["PLB05", "PLB10", "PLB12", "PLB15"]
SA_POOL = [300000, 349999, 350000, 400000, 499999, 500000, 650000, 699999, 700000, 999999, 1000000, 1500000, 3000000, 20000000]
MEB_PLANS = [500, 1000, 2000, 3000, 4000, 5000]

PROFILE_XCU = """<?xml version="1.0" encoding="UTF-8"?>
<oor:items xmlns:oor="http://openoffice.org/2001/registry" xmlns:xs="http://www.w3.org/2001/XMLSchema" xmlns:xsi="http://www.w3.org/2001/XMLSchema-instance">
<item oor:path="/org.openoffice.Office.Calc/Formula/Load"><prop oor:name="OOXMLRecalcMode" oor:op="fuse"><value>0</value></prop></item>
</oor:items>
"""


def random_cases(seed, count):
    rng = random.Random(seed)
    cases = []
    for _ in range(count):
        age = rng.randint(20, 59)
        sa = rng.choice(SA_POOL)
        riders = {}
        if rng.random() < 0.6:
            riders["AP"] = rng.choice([100000, 250000, 500000, 1000000, 2000000, min(5 * sa, 10000000)])
        if rng.random() < 0.5:
            riders["ECARE"] = rng.choice([100000, 300000, 500000, 1000000])
        if rng.random() < 0.5:
            riders["MEB"] = rng.choice(MEB_PLANS)
        cases.append({
            "variant": rng.choice(VARIANTS), "age": age, "sex": rng.choice("MF"),
            "mode": rng.choice(["annual", "semi", "monthly"]), "sumAssured": sa, "riders": riders,
        })
    return cases


def write_inputs(case, path):
    wb = openpyxl.load_workbook(SRC)
    ws = wb["กรอกข้อมูล"]
    ws["F7"] = case["age"]
    ws["F8"] = SEX_TH[case["sex"]]
    ws["F9"] = MODE_TH[case["mode"]]
    ws["D19"] = case["variant"]
    ws["G19"] = case["sumAssured"]
    r = case["riders"]
    ws["G23"] = r.get("AP")
    ws["G24"] = r.get("ECARE")
    ws["G27"] = r.get("MEB", "ไม่ซื้อ")
    wb.save(path)


def read_outputs(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws, cal = wb["กรอกข้อมูล"], wb["Cal"]

    def cell(c):
        v = ws[c].value
        return v if v is not None else 0

    return {
        "modal": {"BASE": cell("I19"), "AP": cell("I23"), "ECARE": cell("I24"), "MEB": cell("I27")},
        "annual": {"BASE": cal["G13"].value or 0, "AP": cal["G17"].value or 0, "ECARE": cal["G18"].value or 0, "MEB": cal["G21"].value or 0},
        "totalModal": cell("I32"),
        "monthlyMessage": ws["J32"].value,
    }


def main():
    if not Path(SOFFICE).exists():
        sys.exit(f"LibreOffice not found at {SOFFICE}. Run: brew reinstall --cask libreoffice")
    spec = json.loads(CASES.read_text(encoding="utf-8"))
    cases = spec["explicit"] + random_cases(spec["random"]["seed"], spec["random"]["count"])

    with tempfile.TemporaryDirectory() as tmpdir:
        tmp = Path(tmpdir)
        profile = tmp / "profile" / "user"
        profile.mkdir(parents=True)
        (profile / "registrymodifications.xcu").write_text(PROFILE_XCU, encoding="utf-8")
        indir, outdir = tmp / "in", tmp / "out"
        indir.mkdir()
        outdir.mkdir()
        for i, c in enumerate(cases):
            write_inputs(c, indir / f"case{i:03}.xlsx")
        cmd = [SOFFICE, f"-env:UserInstallation=file://{tmp / 'profile'}", "--headless",
               "--convert-to", "xlsx", "--outdir", str(outdir)] + sorted(str(p) for p in indir.glob("*.xlsx"))
        subprocess.run(cmd, check=True, capture_output=True)
        results = []
        for i, c in enumerate(cases):
            got = read_outputs(outdir / f"case{i:03}.xlsx")
            results.append({"input": c, "expected": got})
    OUT.write_text(json.dumps({"source": "libreoffice-recalc", "sourceFile": SRC.name, "cases": results},
                              ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"wrote {OUT} with {len(results)} cases")


if __name__ == "__main__":
    main()
