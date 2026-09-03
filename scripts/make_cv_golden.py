#!/usr/bin/env python3
"""Recalculate the ไลฟ์ โพรเทค+ workbook for a few cases and save its own benefit table.

The workbook is the authority on surrender values, so the engine is checked against what
LibreOffice computes here rather than against numbers typed in by hand.
Usage: python3 scripts/make_cv_golden.py
"""
import json
import shutil
import subprocess
import sys
import tempfile
from pathlib import Path

import openpyxl

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "ไฟล์คำนวน" / "Sales proposal_Life Protect Plus 9 19 99_2026821_update benefit graph.xlsx"
OUT = ROOT / "tests" / "golden" / "lifeprotect-cash-values.json"
SEX_TH = {"M": "ชาย", "F": "หญิง"}

CASES = [
    {"age": 35, "sex": "M", "variant": "WLF99H", "sumAssured": 1_000_000},
    {"age": 45, "sex": "F", "variant": "WLF19H", "sumAssured": 500_000},
    {"age": 0, "sex": "F", "variant": "WLF09L", "sumAssured": 300_000},
    {"age": 60, "sex": "M", "variant": "WLF99L", "sumAssured": 2_000_000},
    {"age": 80, "sex": "F", "variant": "WLF19L", "sumAssured": 150_000},
]


def packages():
    rates = json.loads((ROOT / "data" / "rates" / "lifeprotect.json").read_text(encoding="utf-8"))
    return {p["code"]: p for p in rates["base"]["packages"]}


def write_case(path, case, pkgs):
    wb = openpyxl.load_workbook(path)
    ws = wb["กรอกข้อมูล"]
    pkg = pkgs[case["variant"]]
    ws["C5"] = case["age"]
    ws["C6"] = SEX_TH[case["sex"]]
    ws["C7"] = "รายปี"
    ws["C14"] = pkg["name"]
    ws["C15"] = pkg["productName"]
    ws["D20"] = case["sumAssured"]
    wb.save(path)


def read_case(path):
    """Age, policy year and surrender value from both halves of the benefit table."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["ตารางแสดงผลประโยชน์"]
    rows = []
    for age_col, year_col, cv_col, death_col in (("B", "C", "I", "H"), ("K", "L", "R", "Q")):
        for r in range(17, 118):
            age, year, cv = ws[f"{age_col}{r}"].value, ws[f"{year_col}{r}"].value, ws[f"{cv_col}{r}"].value
            if age in (None, "") or year in (None, ""):
                continue
            rows.append({"age": int(age), "policyYear": int(year),
                         "cashValue": int(cv or 0), "deathBenefit": int(ws[f"{death_col}{r}"].value or 0)})
    rows.sort(key=lambda x: x["policyYear"])
    return rows


def main():
    pkgs = packages()
    out = []
    with tempfile.TemporaryDirectory() as tmp:
        tmp = Path(tmp)
        for i, case in enumerate(CASES):
            work = tmp / f"case{i}.xlsx"
            shutil.copy(SRC, work)
            write_case(work, case, pkgs)
            subprocess.run(
                ["soffice", "--headless", "--convert-to", "xlsx", "--outdir", str(tmp / f"out{i}"), str(work)],
                check=True, capture_output=True, timeout=600)
            rows = read_case(tmp / f"out{i}" / f"case{i}.xlsx")
            assert rows, f"case {i} produced no benefit rows"
            out.append({**case, "rows": rows})
            print(f"case {i}: {case['variant']} {case['sex']}{case['age']} -> {len(rows)} rows, "
                  f"last age {rows[-1]['age']} cv {rows[-1]['cashValue']:,}")
    OUT.parent.mkdir(parents=True, exist_ok=True)
    OUT.write_text(json.dumps(out, ensure_ascii=False, indent=1) + "\n", encoding="utf-8")
    print(f"wrote {OUT}")


if __name__ == "__main__":
    sys.exit(main())
